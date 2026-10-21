"""Human-in-the-loop fix application: plain text, conflicts, and Word tracked changes."""
import json
from pathlib import Path

from backend.schemas import (
    Citation, Coverage, Finding, Meta, PolicyDoc, Result, Span, Stats, Suggestion,
)
from backend.tools.apply import apply_fixes, build_docx
from backend.tools.clauses import split_clauses
from backend.tools.parse import parse_document

ROOT = Path(__file__).resolve().parents[2]
CIT = [Citation(doc_id="d", chunk_id="d:0", quote="q", source_title="Regulation", source_url="https://x")]


def mock_result() -> Result:
    return Result.model_validate(json.loads((ROOT / "frontend" / "mock" / "result.json").read_text(encoding="utf-8")))


def test_apply_selected_with_reviewer_edit_keeps_numbering():
    r = mock_result()
    rev = apply_fixes(r, [{"finding_id": "F7", "text": "Records shall be kept for at least five years."}])
    assert len(rev.changes) == 1 and not rev.skipped
    c = rev.changes[0]
    assert rev.full_text[c.start:c.end] == "6.1 Records shall be kept for at least five years."  # label restored
    assert "three years" not in rev.full_text and c.old_text.startswith("6.1 The Bank will retain")
    assert "3.2 Customer identification records" in rev.full_text  # untouched clauses stay


def test_apply_all_resolves_conflicts_and_inserts_missing():
    r = mock_result()
    rev = apply_fixes(r, [{"finding_id": f.id} for f in r.findings])
    assert [s["finding_id"] for s in rev.skipped] == ["F4"]  # F3 and F4 rewrite the same clause; F3 (higher) wins
    assert rev.skipped[0]["conflicts_with"] == "F3"
    ins = next(c for c in rev.changes if c.kind == "insert")
    assert rev.full_text[ins.start:ins.end].startswith("7.2 No director")
    # inserted right after clause 7.1, before section 8
    assert rev.full_text.index("7.1 Staff shall report") < ins.start < rev.full_text.index("8. Training")
    for c in rev.changes:
        assert rev.full_text[c.start:c.end] == c.new_text


def test_empty_suggestion_and_unknown_ids_are_reported():
    r = mock_result()
    r.findings[0].suggestion = Suggestion(text="")
    rev = apply_fixes(r, [{"finding_id": r.findings[0].id}, {"finding_id": "F99"}])
    assert not rev.changes
    assert {s["finding_id"] for s in rev.skipped} == {r.findings[0].id, "F99"}


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _accepted(p) -> str:
    """Paragraph text after accepting all tracked changes (w:t kept, w:delText dropped)."""
    return "".join(t.text or "" for t in p._p.iter(f"{W}t")).strip()


def _deleted_mark(p) -> bool:
    ppr = p._p.find(f"{W}pPr")
    rpr = ppr.find(f"{W}rPr") if ppr is not None else None
    return rpr is not None and rpr.find(f"{W}del") is not None


def _count(docx_path: Path) -> tuple[int, int]:
    import docx

    xml = docx.Document(str(docx_path)).element.body.xml
    return xml.count("<w:ins "), xml.count("<w:del ")


def test_word_export_tracked_option_generated_from_text(tmp_path):
    r = mock_result()
    rev = apply_fixes(r, [{"finding_id": "F7"}, {"finding_id": "F9"}])
    info = build_docx(r, rev, None, tmp_path / "out.docx", tracked=True)
    assert info == {"placed": 2, "unplaced": [], "from_original": False}
    ins, dels = _count(tmp_path / "out.docx")
    assert ins >= 3 and dels >= 1  # replacement (del + ins) and a new paragraph (mark + run)


def _demo_result():
    src = ROOT / "samples" / "demo_digital_bank_aml_policy.docx"
    text = parse_document(src).text
    clauses = split_clauses(text)
    by = {c.label: c for c in clauses}

    def finding(i, label, sug):
        c = by[label]
        return Finding(id=f"F{i}", clause_ids=[c.id], span=Span(start=c.start, end=c.end), gap_type="INSUFFICIENT",
                       severity="high", title="t", explanation="e", citations=CIT, suggestion=Suggestion(text=sug))

    anchor = by["10.2"].end
    findings = [
        finding(1, "11.1", "The Bank will retain records for at least five years."),
        finding(2, "9.1", "9.1 The Bank shall not maintain correspondent relationships with shell banks."),
        Finding(id="F3", clause_ids=[], span=Span(start=anchor, end=anchor), gap_type="MISSING", severity="high",
                title="t", explanation="e", citations=CIT, suggestion=Suggestion(text="10.3 Tipping-off is prohibited.")),
    ]
    r = Result(
        run_id="t", meta=Meta(country="PK", timezone="UTC", as_of="2026-10-17", policy_type="AML/KYC",
                              institution_type="Digital bank", generated_at="x"),
        policy=PolicyDoc(filename=src.name, full_text=text, clauses=clauses), findings=findings, gap_table=[],
        coverage=Coverage(), sources=[], stats=Stats(),
    )
    return r, src


def test_word_export_tracked_option_edits_original_docx(tmp_path):
    r, src = _demo_result()
    rev = apply_fixes(r, [{"finding_id": f.id} for f in r.findings])
    info = build_docx(r, rev, src, tmp_path / "rev.docx", tracked=True)
    assert info["from_original"] and info["unplaced"] == []

    import docx

    d = docx.Document(str(tmp_path / "rev.docx"))
    paras = [_accepted(p) for p in d.paragraphs if _accepted(p) or not _deleted_mark(p)]
    # the text Word shows after "Accept all changes"
    assert "11.1 The Bank will retain records for at least five years." in paras
    i = paras.index("10.2 Alerts will be reviewed by the Compliance team, and suspicious transactions will be reported to the Financial Monitoring Unit within 30 days of the alert being closed.")
    assert paras[i + 1] == "10.3 Tipping-off is prohibited."
    headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "11. Record Keeping" in headings  # original structure and styles preserved
    ins, dels = _count(tmp_path / "rev.docx")
    assert dels == 2 and ins >= 4


def test_word_export_default_is_clean_final_text(tmp_path):
    import docx

    r, src = _demo_result()
    rev = apply_fixes(r, [{"finding_id": f.id} for f in r.findings])
    info = build_docx(r, rev, src, tmp_path / "final.docx")
    assert info["from_original"] and info["unplaced"] == []
    d = docx.Document(str(tmp_path / "final.docx"))
    xml = d.element.body.xml
    assert "<w:ins " not in xml and "<w:del " not in xml and "w:delText" not in xml  # no tracked changes
    paras = [p.text for p in d.paragraphs]
    assert "11.1 The Bank will retain records for at least five years." in paras
    assert not any("four years" in p for p in paras)  # old wording is gone, not struck through
    i = next(k for k, p in enumerate(paras) if p.startswith("10.2 Alerts"))
    assert paras[i + 1] == "10.3 Tipping-off is prohibited."
    fixed = next(p for p in d.paragraphs if p.text.startswith("11.1 The Bank will retain records"))
    assert fixed.style.name == "Normal"
    assert "11. Record Keeping" in [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
