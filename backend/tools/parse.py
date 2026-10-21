"""Parse an uploaded policy (PDF, DOCX, TXT) into a stable ``full_text``.

Normalisation is applied once here; every offset in the system refers to the
returned ``text``.
"""
from __future__ import annotations

import re
from pathlib import Path

from backend.schemas import PageSpan, ParsedDoc

SUPPORTED = {".pdf", ".docx", ".txt"}


class UnsupportedDocument(ValueError):
    pass


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ").replace("﻿", "")
    text = text.replace("­", "")  # soft hyphen
    text = re.sub(r"[\t\f\v ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def _join_pages(pages: list[str]) -> ParsedDoc:
    parts: list[str] = []
    spans: list[PageSpan] = []
    pos = 0
    for i, raw in enumerate(pages, start=1):
        body = normalize_text(raw).rstrip("\n")
        if not body:
            continue
        if parts:
            parts.append("\n\n")
            pos += 2
        spans.append(PageSpan(number=i, start=pos, end=pos + len(body)))
        parts.append(body)
        pos += len(body)
    text = "".join(parts) + "\n"
    return ParsedDoc(text=text, pages=spans)


def parse_pdf_bytes(data: bytes) -> ParsedDoc:
    import fitz  # pymupdf

    with fitz.open(stream=data, filetype="pdf") as doc:
        pages = [page.get_text("text") for page in doc]
    return _join_pages(pages)


def _docx_text(path: Path) -> str:
    import docx  # python-docx

    d = docx.Document(str(path))
    lines: list[str] = []
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(t.text or "" for t in block.iter() if t.tag.endswith("}t"))
            lines.append(text)
        elif tag == "tbl":
            for row in block.iter():
                if row.tag.endswith("}tr"):
                    cells = []
                    for cell in row.iter():
                        if cell.tag.endswith("}tc"):
                            cells.append("".join(t.text or "" for t in cell.iter() if t.tag.endswith("}t")).strip())
                    lines.append(" | ".join(c for c in cells if c))
    return "\n".join(lines)


def parse_document(path: str | Path) -> ParsedDoc:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise UnsupportedDocument(f"Unsupported file type '{ext}'. Upload a PDF, DOCX or TXT file.")
    if ext == ".pdf":
        parsed = parse_pdf_bytes(path.read_bytes())
    elif ext == ".docx":
        text = normalize_text(_docx_text(path))
        parsed = ParsedDoc(text=text, pages=[PageSpan(number=1, start=0, end=len(text))])
    else:
        raw = path.read_bytes()
        encs = ("utf-16",) if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else ()
        decoded = ""
        for enc in (*encs, "utf-8-sig", "cp1252", "latin-1"):
            try:
                decoded = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        text = normalize_text(decoded)
        parsed = ParsedDoc(text=text, pages=[PageSpan(number=1, start=0, end=len(text))])
    if len(parsed.text.strip()) < 20:
        raise UnsupportedDocument("The document appears to be empty or is a scanned image without text.")
    return parsed
