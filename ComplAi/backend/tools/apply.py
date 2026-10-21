"""Apply reviewer-approved fixes to the policy (human-in-the-loop step).

The analysis only *proposes* fixes. A person chooses which ones to accept,
can edit the wording first, and this module turns those decisions into:

* ``apply_fixes``  -> the revised plain text plus a list of changes (for the UI diff)
* ``build_docx``   -> the final Word document: approved fixes are written straight
  into the text (no strikethrough, no highlight). For a Word upload the original
  file is edited in place, keeping its own formatting; other uploads get a
  generated document. ``tracked=True`` instead records each fix as a Word tracked
  change by "ComplAI".

No LLM is involved: applying fixes is deterministic and costs no tokens.
"""
from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from backend.schemas import SEVERITY_RANK, Clause, Finding, Result
from backend.tools.clauses import is_heading
from backend.tools.quote import normalize

AUTHOR = "ComplAI"


@dataclass
class Edit:
    finding_id: str
    kind: str  # "replace" | "insert"
    start: int  # offsets in the ORIGINAL text
    end: int
    new_text: str
    clause_ids: list[str] = field(default_factory=list)
    severity: str = "medium"


@dataclass
class Change:
    finding_id: str
    kind: str
    start: int  # offsets in the REVISED text
    end: int
    old_text: str
    new_text: str


@dataclass
class Revision:
    full_text: str
    changes: list[Change]
    skipped: list[dict]
    edits: list[Edit]


def _clause_range(result: Result, clause_ids: list[str]) -> tuple[int, int, list[Clause]] | None:
    clauses = [c for c in result.policy.clauses if c.id in set(clause_ids)]
    if not clauses:
        return None
    return min(c.start for c in clauses), max(c.end for c in clauses), sorted(clauses, key=lambda c: c.start)


def _with_label(new_text: str, first: Clause) -> str:
    """Keep the clause's own numbering (e.g. '5.3') if the suggestion dropped it."""
    text = new_text.strip()
    label = (first.label or "").strip()
    if label and not re.match(rf"^\(?{re.escape(label.strip('()'))}\)?[.:]?\s", text + " "):
        text = f"{label} {text}"
    return text


def plan_edits(result: Result, selections: list[dict]) -> tuple[list[Edit], list[dict]]:
    """Turn reviewer selections into non-overlapping edits.

    ``selections``: ``[{"finding_id": "F3", "text": "optional reviewer wording"}]``.
    When two selected fixes rewrite the same clause, the higher-severity one is
    applied and the other is reported back so the reviewer can merge them.
    """
    by_id = {f.id: f for f in result.findings}
    edits: list[Edit] = []
    skipped: list[dict] = []
    for sel in selections:
        f: Finding | None = by_id.get(str(sel.get("finding_id", "")))
        if f is None:
            skipped.append({"finding_id": sel.get("finding_id"), "reason": "Unknown finding."})
            continue
        text = (sel.get("text") if sel.get("text") is not None else f.suggestion.text) or ""
        if not text.strip():
            skipped.append({"finding_id": f.id, "reason": "No replacement text to apply."})
            continue
        if f.gap_type == "MISSING" or not f.clause_ids:
            pos = f.span.start
            edits.append(Edit(f.id, "insert", pos, pos, text.strip(), [], f.severity))
            continue
        rng = _clause_range(result, f.clause_ids)
        if rng is None:
            skipped.append({"finding_id": f.id, "reason": "The clause it refers to was not found."})
            continue
        start, end, clauses = rng
        edits.append(Edit(f.id, "replace", start, end, _with_label(text, clauses[0]), [c.id for c in clauses], f.severity))

    # resolve overlapping replacements: higher severity first, then finding order
    def num(fid: str) -> int:
        m = re.search(r"\d+", fid)
        return int(m.group()) if m else 0

    replaces = sorted((e for e in edits if e.kind == "replace"), key=lambda e: (-SEVERITY_RANK.get(e.severity, 1), num(e.finding_id)))
    kept: list[Edit] = []
    for e in replaces:
        clash = next((k for k in kept if e.start < k.end and k.start < e.end), None)
        if clash:
            skipped.append({
                "finding_id": e.finding_id,
                "reason": f"Rewrites the same clause as {clash.finding_id}, which was applied instead. "
                          f"Edit {clash.finding_id}'s wording to include this fix if both are needed.",
                "conflicts_with": clash.finding_id,
            })
        else:
            kept.append(e)
    inserts = [e for e in edits if e.kind == "insert"]
    # an insertion anchored inside a replaced clause moves to that clause's end
    for ins in inserts:
        host = next((k for k in kept if k.start < ins.start < k.end), None)
        if host:
            ins.start = ins.end = host.end
    final = sorted(kept + inserts, key=lambda e: (e.start, 0 if e.kind == "replace" else 1, num(e.finding_id)))
    return final, skipped


def apply_fixes(result: Result, selections: list[dict]) -> Revision:
    text = result.policy.full_text
    edits, skipped = plan_edits(result, selections)
    out: list[str] = []
    changes: list[Change] = []
    cursor = 0
    pos = 0  # length of revised text so far
    for e in edits:
        chunk = text[cursor:e.start]
        out.append(chunk)
        pos += len(chunk)
        if e.kind == "replace":
            new = e.new_text
            changes.append(Change(e.finding_id, "replace", pos, pos + len(new), text[e.start:e.end], new))
            out.append(new)
            pos += len(new)
            cursor = e.end
        else:
            # new clause on its own line after the anchor
            prefix = "\n" if (e.start > 0 and text[e.start - 1:e.start] != "\n") else ""
            suffix = "" if text[e.start:e.start + 1] == "\n" else "\n"
            block = f"{prefix}{e.new_text}{suffix}"
            s = pos + len(prefix)
            changes.append(Change(e.finding_id, "insert", s, s + len(e.new_text), "", e.new_text))
            out.append(block)
            pos += len(block)
            cursor = e.start
    out.append(text[cursor:])
    return Revision(full_text="".join(out), changes=changes, skipped=skipped, edits=edits)


# ---------------------------------------------------------------------------
# Word export with tracked changes
# ---------------------------------------------------------------------------
class _Tracker:
    def __init__(self) -> None:
        self.next_id = 1
        self.date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    def mark(self, tag: str):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        el = OxmlElement(tag)
        el.set(qn("w:id"), str(self.next_id))
        el.set(qn("w:author"), AUTHOR)
        el.set(qn("w:date"), self.date)
        self.next_id += 1
        return el


def _norm(s: str) -> str:
    return normalize(s).strip()


def _delete_runs(paragraph, tr: _Tracker, mark_paragraph: bool) -> None:
    from docx.oxml.ns import qn

    p = paragraph._p
    for r in list(p.findall(qn("w:r"))):
        for t in r.findall(qn("w:t")):
            t.tag = qn("w:delText")
        d = tr.mark("w:del")
        r.addprevious(d)
        d.append(r)
    if mark_paragraph:  # also delete the paragraph mark so the empty line disappears on accept
        rpr = p.get_or_add_pPr().find(qn("w:rPr"))
        if rpr is None:
            from docx.oxml import OxmlElement

            rpr = OxmlElement("w:rPr")
            p.get_or_add_pPr().append(rpr)
        rpr.append(tr.mark("w:del"))


def _inserted_run(text: str, tr: _Tracker, template_run=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ins = tr.mark("w:ins")
    r = OxmlElement("w:r")
    if template_run is not None:
        rpr = template_run.find(qn("w:rPr"))
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    ins.append(r)
    return ins


def _find_paragraphs(paras: list, used: set[int], old_text: str) -> list[int] | None:
    """Indices of consecutive paragraphs whose text makes up ``old_text``."""
    target = _norm(old_text)
    if not target:
        return None
    norms = [_norm(p.text) for p in paras]
    for i, n in enumerate(norms):
        if i in used or not n or not target.startswith(n[: max(1, min(len(n), 60))]):
            continue
        acc = n
        j = i
        while acc != target and len(acc) < len(target) and j + 1 < len(paras):
            j += 1
            if norms[j]:
                acc = f"{acc} {norms[j]}"
        if acc == target or (len(acc) >= len(target) * 0.97 and acc.startswith(target[: int(len(target) * 0.9)])):
            return list(range(i, j + 1))
    return None


def _anchor_paragraph(result: Result, paras: list, used_map: dict[str, list[int]], edit: Edit) -> int | None:
    """Paragraph after which a MISSING clause is inserted: the clause ending at the anchor."""
    before = [c for c in result.policy.clauses if c.end <= edit.start]
    for c in sorted(before, key=lambda c: -c.end):
        if c.id in used_map:
            return used_map[c.id][-1]
        idx = _find_paragraphs(paras, set(), c.text)
        if idx:
            return idx[-1]
    return len(paras) - 1 if paras else None


def _generated_document(full_text: str):
    import docx

    d = docx.Document()
    for line in full_text.split("\n"):
        if not line.strip():
            continue
        if is_heading(line):
            d.add_heading(line.strip(), level=2)
        else:
            d.add_paragraph(line.strip())
    return d


def _plain_run(text: str, template_run=None):
    """A normal run carrying the original paragraph's character formatting."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r = OxmlElement("w:r")
    if template_run is not None:
        rpr = template_run.find(qn("w:rPr"))
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


def _clear_content(paragraph) -> None:
    """Remove everything in a paragraph except its paragraph properties."""
    from docx.oxml.ns import qn

    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def build_docx(result: Result, revision: Revision, source: Path | None, out_path: Path, tracked: bool = False) -> dict:
    """Write the revised policy as .docx. Clean final text by default; ``tracked``
    records each fix as a Word tracked change instead. Returns placed / unplaced counts."""
    import docx

    use_source = source is not None and source.suffix.lower() == ".docx" and source.exists()
    d = docx.Document(str(source)) if use_source else _generated_document(result.policy.full_text)
    paras = list(d.paragraphs)
    tr = _Tracker()
    used: set[int] = set()
    clause_paras: dict[str, list[int]] = {}
    unplaced: list[Edit] = []
    old_text_of = {e.finding_id: result.policy.full_text[e.start:e.end] for e in revision.edits}

    # replacements
    for e in revision.edits:
        if e.kind != "replace":
            continue
        idx = _find_paragraphs(paras, used, old_text_of[e.finding_id])
        if not idx:
            unplaced.append(e)
            continue
        used.update(idx)
        first = paras[idx[0]]
        template = copy.deepcopy(first._p.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"))
        if tracked:
            for cid in e.clause_ids:
                clause_paras[cid] = idx
            for k, pi in enumerate(idx):
                _delete_runs(paras[pi], tr, mark_paragraph=k > 0)
            first._p.append(_inserted_run(e.new_text, tr, template))
        else:
            for cid in e.clause_ids:
                clause_paras[cid] = idx[:1]
            _clear_content(first)
            first._p.append(_plain_run(e.new_text, template))
            for pi in idx[1:]:  # a clause that spanned several paragraphs collapses into one
                el = paras[pi]._p
                el.getparent().remove(el)

    # insertions (new clauses for MISSING findings), in reverse so order is preserved
    for e in sorted((e for e in revision.edits if e.kind == "insert"), key=lambda e: -e.start):
        anchor = _anchor_paragraph(result, paras, clause_paras, e)
        if anchor is None:
            unplaced.append(e)
            continue
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        ap = paras[anchor]._p
        newp = OxmlElement("w:p")
        ppr = ap.find(qn("w:pPr"))
        if ppr is not None:
            newp.append(copy.deepcopy(ppr))
        ppr = newp.get_or_add_pPr()
        rpr = ppr.find(qn("w:rPr"))
        if rpr is not None:
            for old in list(rpr.findall(qn("w:del"))) + list(rpr.findall(qn("w:ins"))):
                rpr.remove(old)
        template = ap.find(qn("w:r"))
        if tracked:
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                ppr.append(rpr)
            rpr.append(tr.mark("w:ins"))  # the new paragraph mark itself is an insertion
            newp.append(_inserted_run(e.new_text, tr, template))
        else:
            newp.append(_plain_run(e.new_text, template))
        ap.addnext(newp)

    if unplaced:  # e.g. clauses inside tables: list them at the end rather than drop them
        d.add_heading("Additional clauses", level=2)
        for e in unplaced:
            p = d.add_paragraph()
            p._p.append(_inserted_run(f"[{e.finding_id}] {e.new_text}", tr) if tracked else _plain_run(e.new_text))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out_path))
    return {"placed": len(revision.edits) - len(unplaced), "unplaced": [e.finding_id for e in unplaced],
            "from_original": use_source}
